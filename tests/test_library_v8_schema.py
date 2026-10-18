"""Synthetic modern three-table contracts; no user corpus is stored here."""
import csv
import io
import json
import sqlite3
import zipfile

import pytest
from openpyxl import Workbook

from inkreference.csvio import export_legacy_zip, parse_files, preview_import, commit_import
from inkreference.database import DataError, DataStore, TABLES, validate_values
from inkreference.library_exchange import parse_library_files, preview_library_import, commit_library_import, export_library


def j(value):
    return json.dumps(value,ensure_ascii=False,separators=(",",":"))


def table_bytes(table, rows):
    stream = io.StringIO(newline="")
    writer = csv.DictWriter(stream,fieldnames=TABLES[table]["columns"])
    writer.writeheader(); writer.writerows([validate_values(table,row) for row in rows])
    return stream.getvalue().encode("utf-8-sig")


def sources():
    works = [
        {"作品ID":"W1","作品":"同名","类别":"小说","发布日期":"2020","用户是否完整体验过":"是","热度权重":"3","关联作品":j([{"作品ID":"W2","关联方式":"改编"}]),"关联角色":j(["P1"]),"关联cliché":j(["C1"])},
        {"作品ID":"W2","作品":"同名","类别":"影视","用户是否完整体验过":"否","关联角色":j(["P1"]),"关联cliché":j(["C1"])},
        {"作品ID":"W3","作品":"第三部","类别":"小说","用户是否完整体验过":"未确认"},
    ]
    characters = [{"人物ID":"P1","角色":"某人","别名":j(["旧名","称号"]),"后续能力":"能力😀\r\n变化","关联作品":j([{"作品ID":"W1"},{"作品ID":"W2","人物设定":{"开场身份":"学生","后续能力":"第一章"},"自定义":{"章节":1}}]),"关联cliché":j([{"clichéID":"C1","作品ID":"W1"},{"clichéID":"C1","作品ID":"W2"}])}]
    cliches = [{"clichéID":"C1","cliché内容":"独立内容，不需要前缀","cliché程度（%）":"3.14","命中作品数":"2","命中热度合计":"7","关联作品":j([{"作品ID":"W1","作品名":"同名","关联内容":"原版细节"},{"作品ID":"W2","作品名":"同名","关联内容":"改编细节"}]),"关联角色":j([{"人物ID":"P1","作品ID":"W1"},{"人物ID":"P1","作品ID":"W2"}])}]
    return [("作品.csv",table_bytes("works",works)),("角色.csv",table_bytes("characters",characters)),("cliche.csv",table_bytes("cliches",cliches))]


def imported(tmp_path):
    store = DataStore(tmp_path/"source.db")
    batch = parse_library_files(sources())
    assert not batch.errors
    assert not preview_library_import(store,batch)["errors"]
    assert commit_library_import(store,batch)["imported"] == 5
    return store


@pytest.mark.parametrize("fmt",["db","csv","xlsx"])
def test_three_table_exact_roundtrip_multiverse_metadata(tmp_path,fmt):
    store = imported(tmp_path)
    ext = "zip" if fmt == "csv" else fmt
    data = export_library(store,fmt)
    batch = parse_library_files([("all."+ext,data)])
    assert not batch.errors
    target = DataStore(tmp_path/(fmt+".db"))
    assert not preview_library_import(target,batch)["errors"]
    commit_library_import(target,batch)
    for table in TABLES:
        assert store.all_records(table) == target.all_records(table)
    character = target.get_record("characters","P1")
    assert json.loads(character["values"]["关联作品"])[1]["人物设定"]["开场身份"] == "学生"
    if fmt == "csv":
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            assert set(archive.namelist()) == {"works.csv","characters.csv","cliches.csv","manifest.json"}
    rev = character["revision"]
    assert not preview_library_import(target,batch)["errors"]
    assert commit_library_import(target,batch,"update")["skipped"] == 5
    assert target.get_record("characters","P1")["revision"] == rev


def test_edit_seen_relation_metadata_and_delete_work_keeps_shared_character(tmp_path):
    store = imported(tmp_path)
    work = store.get_record("works","W1")
    work = store.update_record("works","W1",{"看过":"否","作品":"新名"},work["revision"])
    assert work["values"]["用户是否完整体验过"] == "否"
    assert json.loads(store.get_record("cliches","C1")["values"]["关联作品"])[0]["作品名"] == "新名"
    with pytest.raises(DataError,match="已被修改"):
        store.update_record("works","W1",{"看过":"是"},work["revision"]-1)
    with pytest.raises(DataError,match="关联"):
        store.delete_record("works","W1",work["revision"])
    assert store.delete_record("works","W1",work["revision"],cascade=True) == {"deleted":1}
    person = store.get_record("characters","P1")
    assert json.loads(person["values"]["关联作品"]) == [{"作品ID":"W2","人物设定":{"开场身份":"学生","后续能力":"第一章"},"自定义":{"章节":1}}]
    assert all(item.get("作品ID") != "W1" for item in json.loads(person["values"]["关联cliché"]))
    assert store.get_record("cliches","C1")["values"]["命中热度合计"] == "7"
    with store.connect() as connection:
        assert not connection.execute("PRAGMA foreign_key_check").fetchall()


def test_sql_filters_before_paging_full_facets_and_empty_work_ids(tmp_path):
    store = imported(tmp_path)
    result = store.list_records("works",limit=1,filters={"category":"小说","seen":"unknown"},sort_by="作品",sort_direction="desc")
    assert result["total"] == 1 and result["items"][0]["id"] == "W3"
    assert result["facets"]["categories"] == ["小说","影视"]
    assert store.list_records("works",filters={"work_ids":[]})["total"] == 0
    assert [item["id"] for item in store.list_records("works",filters={"work_ids":["W2"]},limit=1)["items"]] == ["W2"]
    assert store.list_records("characters",work_id="W2")["total"] == 1
    assert store.list_records("works",q="同名")["total"] == 2


def test_missing_cross_version_and_reverse_claims_are_atomic(tmp_path):
    files = sources()
    batch = parse_library_files(files[:2])
    assert "cliches.csv" in " ".join(batch.errors)
    files[1] = ("角色.csv",files[1][1].replace(b"W2",b"MISSING"))
    batch = parse_library_files(files)
    assert "角色.csv" in " ".join(batch.errors) and "作品.csv" in " ".join(batch.errors) and "MISSING" in " ".join(batch.errors)
    store = DataStore(tmp_path/"target.db")
    assert preview_library_import(store,batch)["errors"]
    with pytest.raises(DataError): commit_library_import(store,batch)
    assert store.stats()["works"] == 0


def test_manual_plain_xlsx_three_sheets_is_supported(tmp_path):
    book = Workbook();book.remove(book.active)
    for (name,data),table in zip(sources(),TABLES):
        sheet = book.create_sheet(name[:-4]);reader = csv.reader(io.StringIO(data.decode("utf-8-sig")))
        for row in reader: sheet.append(row)
    output = io.BytesIO();book.save(output)
    batch = parse_library_files([("manual.xlsx",output.getvalue())])
    assert not batch.errors
    store = DataStore(tmp_path/"xlsx.db")
    assert not preview_library_import(store,batch)["errors"]
    commit_library_import(store,batch)
    assert store.stats()["characters"] == 1


def test_schema3_migration_archives_and_explicit_v8_export(tmp_path):
    from pathlib import Path
    fixture_dir = Path(__file__).parent/"fixtures"
    store = DataStore(tmp_path/"old.db")
    # Exercise CRLF explicitly on every host, including embedded CSV newlines.
    old = [(table+".csv", (fixture_dir/(table+".csv")).read_bytes().replace(b"\r\n", b"\n").replace(b"\n", b"\r\n")) for table in TABLES]
    input_bytes = dict(old)
    batch = parse_files(old);assert not preview_import(store,batch)["errors"];commit_import(store,batch)
    # Simulate a persisted v3 database: v4 tables did not exist in this state.
    with store.connect() as connection:
        connection.execute("DELETE FROM schema_migrations WHERE version>=4")
        connection.execute("DROP TABLE library_links");connection.execute("DROP TABLE legacy_values")
    migrated = DataStore(tmp_path/"old.db")
    assert "剧情补充说明" not in migrated.get_record("works","W950")["values"]
    raw = export_legacy_zip(migrated)
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        for table,name in [("works","works_library_v8.csv"),("characters","character_database_v8.csv"),("cliches","cliche_database_v8.csv")]:
            expected = list(csv.DictReader(io.StringIO(input_bytes[table+".csv"].decode("utf-8-sig"), newline="")))
            actual = list(csv.DictReader(io.StringIO(archive.read(name).decode("utf-8-sig"), newline="")))
            assert actual == expected
    target = DataStore(tmp_path/"fresh.db")
    new = parse_library_files([("all.db",export_library(migrated,"db"))])
    assert not preview_library_import(target,new)["errors"];commit_library_import(target,new)
    with zipfile.ZipFile(io.BytesIO(export_legacy_zip(target))) as actual, zipfile.ZipFile(io.BytesIO(raw)) as expected:
        assert {name:actual.read(name) for name in actual.namelist()} == {name:expected.read(name) for name in expected.namelist()}


def test_seen_toggle_does_not_rebuild_relation_index(tmp_path,monkeypatch):
    store = imported(tmp_path)
    work = store.get_record("works","W1")
    def unexpected(*args): raise AssertionError("ordinary edits must not rebuild the full link index")
    monkeypatch.setattr(store,"_rebuild_library_links",unexpected)
    updated = store.update_record("works","W1",{"用户是否完整体验过":"否"},work["revision"])
    assert updated["values"]["用户是否完整体验过"] == "否"


def test_work_relation_direction_and_character_context_edits(tmp_path):
    store = imported(tmp_path)
    work = store.get_record("works","W3")
    store.update_record("works","W3",{"关联作品":j([{"作品ID":"W2","关联方式":"续集"}])},work["revision"])
    assert json.loads(store.get_record("works","W2")["values"]["关联作品"]) == []
    person = store.get_record("characters","P1")
    store.update_record("characters","P1",{"关联cliché":j([{"clichéID":"C1","作品ID":"W2"}])},person["revision"])
    assert json.loads(store.get_record("cliches","C1")["values"]["关联角色"]) == [{"人物ID":"P1","作品ID":"W2"}]


def test_migration_preserves_current_legacy_relation_views_after_edits(tmp_path):
    store = DataStore(tmp_path/"edited-old.db")
    store.create_record("works",{"作品ID":"W1","作品":"旧作","版本与原作关系":"小说"})
    store.create_record("cliches",{"cliché内容":"C1｜发现"})
    item = store.get_record("works","W1")
    store.update_record("works","W1",{"典型桥段（剧透）":"C1：当前关联细节"},item["revision"])
    before = {table:[record["values"] for record in store.all_records(table)] for table in TABLES}
    with store.transaction() as connection:
        store._upgrade_legacy_rows(connection);store._rebuild_library_links(connection)
    with zipfile.ZipFile(io.BytesIO(export_legacy_zip(store))) as archive:
        for table,name in [("works","works_library_v8.csv"),("characters","character_database_v8.csv"),("cliches","cliche_database_v8.csv")]:
            assert list(csv.DictReader(io.StringIO(archive.read(name).decode("utf-8-sig")))) == before[table]
    batch = parse_library_files([("legacy.zip",export_legacy_zip(store))])
    assert not batch.errors


def test_character_cliche_context_reverse_missing_and_distinct_descriptions(tmp_path):
    files = sources()
    character_rows = list(csv.DictReader(io.StringIO(files[1][1].decode("utf-8-sig"))))
    cliche_rows = list(csv.DictReader(io.StringIO(files[2][1].decode("utf-8-sig"))))
    character_rows[0]["关联cliché"] = j([{"clichéID":"C1","作品ID":"W1","关联内容":"版本一"},{"clichéID":"C1","作品ID":"W2","关联内容":"版本二"}])
    cliche_rows[0]["关联角色"] = j([{"人物ID":"P1","作品ID":"W1","关联内容":"版本一"},{"人物ID":"P1","作品ID":"W2","关联内容":"版本二"}])
    files[1] = ("角色.csv",table_bytes("characters",character_rows));files[2] = ("cliche.csv",table_bytes("cliches",cliche_rows))
    valid = parse_library_files(files)
    assert not valid.errors
    store = DataStore(tmp_path/"context.db")
    assert not preview_library_import(store,valid)["errors"];commit_library_import(store,valid)
    cliche_rows[0]["关联角色"] = j([{"人物ID":"P1","作品ID":"W1","关联内容":"版本一"}])
    files[2] = ("cliche.csv",table_bytes("cliches",cliche_rows))
    invalid = parse_library_files(files)
    errors = " ".join(invalid.errors)
    assert all(label in errors for label in ("角色.csv","cliche.csv","P1","C1","W2"))
    # Also enforce this if a trusted internal caller bypasses the file parser.
    with pytest.raises(DataError,match="W2"):
        with store.transaction() as connection:
            record = store.get_record("cliches","C1")
            store._update_raw(connection,"cliches","C1",cliche_rows[0],record["revision"])
            store._rebuild_library_links(connection)
    assert len(json.loads(store.get_record("cliches","C1")["values"]["关联角色"])) == 2


def test_manual_xlsx_duplicate_headers_never_overwrite_ids(tmp_path):
    book = Workbook();book.remove(book.active)
    for table in TABLES:
        sheet = book.create_sheet(table)
        columns = list(TABLES[table]["columns"])
        if table == "works":
            sheet.append([*columns,"作品ID"])
            sheet.append(["W-FIRST",*([""]*(len(columns)-1)),"W-LAST"])
        else: sheet.append(columns)
    output = io.BytesIO();book.save(output)
    batch = parse_library_files([("duplicates.xlsx",output.getvalue())])
    assert "重复字段" in " ".join(batch.errors)
    assert not batch.tables
