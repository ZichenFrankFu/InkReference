"""Single-choice cliche classification and lossless schema-3/4 upgrades."""
from __future__ import annotations

import csv
import io
import json
from pathlib import Path

import pytest

from inkreference.csvio import ImportBatch, ImportTable, commit_import, parse_files, preview_import
from inkreference.database import (
    CLICHE_CHARACTER_TRAITS, CLICHE_PLOT_MECHANISMS, CLICHE_TAGS,
    DataError, DataStore, LEGACY_TABLES, SCHEMA4_COLUMNS, TABLES, validate_values,
)
from inkreference.library_exchange import (
    FORMAT, FORMAT_VERSION, LEGACY_COLUMNS, SCHEMA4_COLUMNS as EXCHANGE_SCHEMA4_COLUMNS,
    COLUMNS, _digest, _export_csv, _export_db, _export_xlsx,
    commit_library_import, export_library, parse_library_files, preview_library_import,
)


def csv_bytes(columns, rows):
    stream = io.StringIO(newline="")
    writer = csv.DictWriter(stream, fieldnames=columns)
    writer.writeheader()
    writer.writerows(rows)
    return stream.getvalue().encode("utf-8-sig")


def modern_rows():
    return {
        "works": [validate_values("works", {"作品ID":"W1", "作品":"测试书", "类别":"小说"})],
        "characters": [validate_values("characters", {"人物ID":"P1", "角色":"甲", "角色定位":"主角"})],
        "cliches": [validate_values("cliches", {"clichéID":"C1", "cliché内容":"原始\r\n内容😀"})],
    }


def old_rows(version):
    if version == 3:
        fixtures = Path(__file__).parent / "fixtures"
        return {table:list(csv.DictReader(io.StringIO((fixtures/(table+".csv")).read_bytes().decode("utf-8-sig"), newline=""))) for table in TABLES}
    return {table:[{column:row[column] for column in SCHEMA4_COLUMNS[table]} for row in rows] for table,rows in modern_rows().items()}


def portable(rows, version, fmt):
    columns = {3:LEGACY_COLUMNS, 4:EXCHANGE_SCHEMA4_COLUMNS, 5:COLUMNS}[version]
    manifest = {
        "format":FORMAT, "format_version":FORMAT_VERSION,
        "schema_version":version, "bundle_id":"synthetic-old-library",
        "tables":{table:{"columns":columns[table], "rows":len(values), "sha256":_digest(values)} for table,values in rows.items()},
    }
    raw = {"db":_export_db,"csv":_export_csv,"xlsx":_export_xlsx}[fmt](rows,manifest)
    return [("library."+("zip" if fmt == "csv" else fmt), raw)]


def import_files(store, files, policy="skip"):
    batch = parse_library_files(files)
    assert not batch.errors, batch.errors
    preview = preview_library_import(store,batch)
    assert not preview["errors"], preview
    result = commit_library_import(store,batch,policy)
    return batch, preview, result


def test_two_independent_single_choices_clear_and_revision_guard(tmp_path):
    store = DataStore(tmp_path/"tags.db")
    item = store.create_record("cliches", {"clichéID":"C1", "cliché内容":"选择", "人设塑造":"强大", "剧情机制":"信息反转"})
    assert item["values"]["人设塑造"] == "强大"
    assert item["values"]["剧情机制"] == "信息反转"
    previous = item
    item = store.update_record("cliches","C1",{"人设塑造":"共情"},item["revision"])
    assert item["values"]["剧情机制"] == "信息反转"
    assert item["revision"] == previous["revision"] + 1
    with pytest.raises(DataError,match="已被修改"):
        store.update_record("cliches","C1",{"剧情机制":"代价交换"},previous["revision"])
    assert store.get_record("cliches","C1") == item
    item = store.update_record("cliches","C1",{"人设塑造":"","剧情机制":""},item["revision"])
    assert all(item["values"][field] == "" for field in CLICHE_TAGS)


@pytest.mark.parametrize("field,value", [
    ("人设塑造","强大、智谋"), ("人设塑造",'["强大"]'),
    ("人设塑造","信息反转"), ("剧情机制","强大"),
    ("剧情机制","信息反转,代价交换"), ("剧情机制","未知"),
    ("人设塑造",["强大"]), ("剧情机制",None),
])
def test_invalid_classification_is_atomic(tmp_path,field,value):
    store = DataStore(tmp_path/"invalid.db")
    item = store.create_record("cliches",{"clichéID":"C1","cliché内容":"原文"})
    with pytest.raises(DataError):
        store.update_record("cliches","C1",{"cliché内容":"不应保存",field:value},item["revision"])
    assert store.get_record("cliches","C1") == item
    with pytest.raises(DataError):
        store.create_record("cliches",{"clichéID":"C2",field:value})
    assert store.stats()["cliches"] == 1


def test_import_bypass_validation_rolls_back_earlier_rows(tmp_path):
    store = DataStore(tmp_path/"rollback.db")
    batch = ImportBatch(tables=[ImportTable("cliches","synthetic.csv",[
        validate_values("cliches",{"clichéID":"C1","人设塑造":"强大"}),
        {**validate_values("cliches",{"clichéID":"C2"}),"剧情机制":"信息反转、机械降神"},
    ])])
    result = preview_import(store,batch)
    assert result["errors"]
    assert store.stats()["cliches"] == 0
    with pytest.raises(DataError): commit_import(store,batch,"update")
    assert store.stats()["cliches"] == 0


def test_classification_filters_apply_before_paging_and_facets_are_global(tmp_path):
    store = DataStore(tmp_path/"filters.db")
    store.create_record("works",{"作品ID":"W1","作品":"第一部","类别":"小说"})
    store.create_record("works",{"作品ID":"W2","作品":"第二部","类别":"影视"})
    for index,role in enumerate(("主角","配角","主角"),1):
        store.create_record("characters",{"人物ID":f"P{index}","角色":f"人物{index}","角色定位":role})
    cases = [("强大","信息反转"),("共情","信息反转"),("强大","代价交换"),("强大","信息反转")]
    for index,(trait,mechanism) in enumerate(cases,1):
        store.create_record("cliches",{"clichéID":f"C{index}","cliché内容":f"桥段{index}","人设塑造":trait,"剧情机制":mechanism})
    result = store.list_records("cliches",filters={"character_trait":"强大","plot_mechanism":"信息反转"},limit=1,offset=1,sort_by="id")
    assert result["total"] == 2
    assert [row["id"] for row in result["items"]] == ["C4"]
    assert result["facets"]["character_traits"] == list(CLICHE_CHARACTER_TRAITS)
    assert result["facets"]["plot_mechanisms"] == list(CLICHE_PLOT_MECHANISMS)
    assert result["facets"]["roles"] == ["主角","配角"]
    assert result["facets"]["works"] == [{"id":"W1","title":"第一部"},{"id":"W2","title":"第二部"}]
    assert store.list_records("characters",filters={"role":"主角"},limit=1,offset=1)["items"][0]["id"] == "P3"
    assert store.list_records("cliches",filters={"人设塑造":"共情"})["total"] == 1
    with pytest.raises(DataError,match="筛选值无效"):
        store.list_records("cliches",filters={"plot_mechanism":"A,B"})
    with pytest.raises(DataError,match="不支持"):
        store.list_records("characters",filters={"character_trait":"强大"})


@pytest.mark.parametrize("version",[3,4])
@pytest.mark.parametrize("fmt",["db","csv","xlsx"])
def test_old_portable_formats_upgrade_without_false_digest_warning(tmp_path,version,fmt):
    rows = old_rows(version)
    store = DataStore(tmp_path/"upgraded.db")
    batch,preview,_ = import_files(store,portable(rows,version,fmt))
    assert not [message for message in preview["warnings"] if "导出时不同" in message]
    assert store.meta()["schema_version"] == 5
    for record in store.all_records("cliches"):
        assert record["values"]["人设塑造"] == record["values"]["剧情机制"] == ""
    if version == 4:
        for table,source in rows.items():
            actual = store.all_records(table)
            assert [{key:item["values"][key] for key in SCHEMA4_COLUMNS[table]} for item in actual] == source
        _,_,repeat = import_files(store,portable(rows,version,fmt),"update")
        assert repeat["updated"] == 0 and repeat["skipped"] == sum(map(len,rows.values()))
    _,_,restored = import_files(DataStore(tmp_path/"restore.db"),[("new.db",export_library(store,"db"))])
    assert restored["imported"] == sum(len(store.all_records(table)) for table in TABLES)
    if version == 3:
        with store.connect() as connection:
            archived = {(row["table_key"],row["record_id"]):json.loads(row["values_json"]) for row in connection.execute("SELECT * FROM legacy_values")}
        assert archived and all("人设塑造" not in values for values in archived.values())


def test_schema4_raw_csv_plain_xlsx_empty_tables_and_legacy_parser(tmp_path):
    from openpyxl import Workbook
    rows = old_rows(4)
    files = [(table+".csv",csv_bytes(SCHEMA4_COLUMNS[table],values)) for table,values in rows.items()]
    batch = parse_files(files)
    assert not batch.errors
    assert batch.tables[-1].rows[0]["人设塑造"] == ""
    source = DataStore(tmp_path/"plain.db")
    import_files(source,files)
    for empty in (False,True):
        book = Workbook();book.remove(book.active)
        for table,columns in SCHEMA4_COLUMNS.items():
            sheet = book.create_sheet(table);sheet.append(columns)
            if not empty:
                for row in rows[table]:sheet.append([row[column] for column in columns])
        output = io.BytesIO();book.save(output)
        target = DataStore(tmp_path/f"plain-{empty}.db")
        import_files(target,[("plain.xlsx",output.getvalue())])
        assert target.stats()["cliches"] == (0 if empty else 1)
    emptyfiles = [(table+".csv",csv_bytes(columns,[])) for table,columns in SCHEMA4_COLUMNS.items()]
    target = DataStore(tmp_path/"empty.db")
    import_files(target,emptyfiles)
    assert target.stats()["cliches"] == 0


@pytest.mark.parametrize("fmt",["db","csv","xlsx"])
def test_schema5_classifications_roundtrip(tmp_path,fmt):
    source = DataStore(tmp_path/"source.db")
    for table,rows in modern_rows().items():
        for row in rows: source.create_record(table,row)
    item = source.get_record("cliches","C1")
    source.update_record("cliches","C1",{"人设塑造":"宿命","剧情机制":"顿悟突破"},item["revision"])
    target = DataStore(tmp_path/"target.db")
    _,preview,_ = import_files(target,[("all."+("zip" if fmt == "csv" else fmt),export_library(source,fmt))])
    assert not preview["warnings"]
    assert target.get_record("cliches","C1")["values"] == source.get_record("cliches","C1")["values"]


def test_schema4_persisted_database_upgrade_preserves_revisions_and_original_fields(tmp_path):
    path = tmp_path/"migrated.db"
    source = DataStore(path)
    item = source.create_record("cliches",{"clichéID":"C1","cliché内容":"原字\r\n😀"})
    with source.transaction() as connection:
        connection.execute("DELETE FROM schema_migrations WHERE version=5")
        connection.execute("UPDATE cliches SET values_json=? WHERE id='C1'",(json.dumps({key:item["values"][key] for key in SCHEMA4_COLUMNS["cliches"]},ensure_ascii=False),))
    migrated = DataStore(path)
    result = migrated.get_record("cliches","C1")
    assert result["revision"] == item["revision"]
    assert result["values"] == item["values"]
    with migrated.connect() as connection:
        assert [row[0] for row in connection.execute("SELECT version FROM schema_migrations ORDER BY version")] == [1,2,3,4,5]
    assert DataStore(path).get_record("cliches","C1") == result
