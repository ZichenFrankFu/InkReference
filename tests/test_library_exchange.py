"""The three portable formats carry one complete, reviewable material library."""
from __future__ import annotations

import csv
import io
import sqlite3
import zipfile
from pathlib import Path

import pytest
from openpyxl import load_workbook

from inkreference.database import TABLES, DataError, DataStore
from inkreference.library_exchange import (
    COLUMNS, commit_library_import, export_library,
    library_formats, parse_library_files, preview_library_import,
)

FIXTURES = Path(__file__).parent / "fixtures"


def fixtures():
    return [(f"{table}_v8.csv", (FIXTURES / f"{table}.csv").read_bytes()) for table in TABLES]


@pytest.fixture
def populated(tmp_path):
    store = DataStore(tmp_path / "source.db")
    batch = parse_library_files(fixtures())
    assert preview_library_import(store, batch)["errors"] == []
    commit_library_import(store, batch)
    store.create_material_relation({"id": "R001", "source_table": "characters", "source_id": "CH-W900-01", "target_table": "works", "target_id": "W902", "relation_type": "借鉴对象", "description": "以电影版为对照。\n保留引号\"与日期 2020-03。", "metadata_json": '{"出处":"手动整理","confidence":0.5}'})
    return store


def roundtrip(tmp_path, store, fmt):
    data = export_library(store, fmt)
    ext = "zip" if fmt == "csv" else fmt
    batch = parse_library_files([(f"materials.{ext}", data)])
    target = DataStore(tmp_path / f"restored-{fmt}.db")
    result = preview_library_import(target, batch)
    assert result["errors"] == [], result
    assert not result["warnings"], result["warnings"]
    commit_library_import(target, batch)
    return target, data


def raw_relations(store):
    with store.connect() as connection:
        return [{key: row[key] for key in COLUMNS["material_relations"]} for row in connection.execute("SELECT * FROM material_relations ORDER BY id")]


def unpack_csv(data):
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        return [(name, archive.read(name)) for name in archive.namelist()]


def csv_data(columns, rows):
    output = io.StringIO(newline="")
    writer = csv.DictWriter(output, fieldnames=columns)
    writer.writeheader()
    writer.writerows(rows)
    return output.getvalue().encode("utf-8-sig")


@pytest.mark.parametrize("fmt", ["db", "csv", "xlsx"])
def test_complete_semantic_roundtrip_all_formats(tmp_path, populated, fmt):
    target, raw = roundtrip(tmp_path, populated, fmt)
    for table in TABLES:
        assert [row["values"] for row in target.all_records(table)] == [row["values"] for row in populated.all_records(table)]
    assert raw_relations(target) == raw_relations(populated)
    assert target.stats() == populated.stats()
    if fmt == "db":
        assert raw.startswith(b"SQLite format 3\x00")
    if fmt == "csv":
        names = {name for name, _ in unpack_csv(raw)}
        assert names == {"works.csv","characters.csv","cliches.csv","material_relations.csv","legacy_values.csv","catalog_field_archive.csv","manifest.json"}
        assert "character_contexts.csv" not in names


@pytest.mark.parametrize("fmt", ["db", "csv", "xlsx"])
def test_empty_library_exports_keep_every_table(tmp_path, fmt):
    target, _ = roundtrip(tmp_path, DataStore(tmp_path / "empty.db"), fmt)
    assert target.stats()["works"] == 0


@pytest.mark.parametrize("fmt", ["db", "csv", "xlsx"])
def test_literal_formula_dates_unknowns_newlines_and_non_bmp(tmp_path, populated, fmt):
    row = populated.get_record("works", "W950")
    text = '=SUM(A1:A2)\r\n𰻞emoji😀\n引号"，未知与【待确认】\r仅回车'
    populated.update_record("works", row["id"], {"剧情梗概": text, "发布日期": "2020-03", "热度权重": "00000012345678901234567890"}, row["revision"])
    target, _ = roundtrip(tmp_path, populated, fmt)
    result = target.get_record("works", "W950")["values"]
    assert result["剧情梗概"] == text
    assert result["发布日期"] == "2020-03"
    assert result["热度权重"] == "00000012345678901234567890"


def test_xlsx_long_and_xml_control_text_are_lossless(tmp_path, populated):
    row = populated.get_record("works", "W950")
    text = ("长文本😀\n" * 10000) + "\x08不可见控制字符"
    populated.update_record("works", row["id"], {"剧情梗概": text}, row["revision"])
    target, raw = roundtrip(tmp_path, populated, "xlsx")
    assert target.get_record("works", "W950")["values"]["剧情梗概"] == text
    book = load_workbook(io.BytesIO(raw), read_only=True)
    assert book["_long_text"].sheet_state == "hidden"
    assert book["_manifest"].sheet_state == "hidden"
    book.close()


def test_xlsx_formula_exported_as_text_and_import_formula_rejected(populated):
    row = populated.get_record("works", "W950")
    populated.update_record("works", row["id"], {"剧情梗概": "=1+1"}, row["revision"])
    raw = export_library(populated, "xlsx")
    book = load_workbook(io.BytesIO(raw))
    cell = book["works"]["G4"]
    assert cell.value == "=1+1" and cell.data_type == "s"
    cell.value = "=2+2"
    output = io.BytesIO()
    book.save(output)
    batch = parse_library_files([("edited.xlsx", output.getvalue())])
    assert "works!G4" in " ".join(batch.errors)
    assert "公式" in " ".join(batch.errors)


def test_missing_legacy_tables_name_each_required_file(tmp_path):
    batch = parse_library_files([fixtures()[0]])
    result = preview_library_import(DataStore(tmp_path / "a.db"), batch)
    assert "character_database_v8.csv" in " ".join(result["errors"])
    assert "cliche_database_v8.csv" in " ".join(result["errors"])
    assert len(result["required_tables"]) == 3
    assert not result["required_tables"][1]["present"]


def test_legacy_filename_version_pair_is_actionable():
    files = fixtures()
    files[1] = ("character_database_v7.csv", files[1][1])
    batch = parse_library_files(files)
    errors = " ".join(batch.errors)
    assert "character_database_v7.csv" in errors and "works_v8.csv" in errors
    assert "版本疑似不一致" in errors


def test_missing_parent_crossfile_names_and_id():
    files = fixtures()
    files[1] = (files[1][0], files[1][1].replace(b"W900", b"W999"))
    errors = " ".join(parse_library_files(files).errors)
    assert "characters_v8.csv" in errors and "works_v8.csv" in errors and "W999" in errors


def test_header_diagnostic_names_missing_column():
    files = fixtures()
    files[0] = ("works_v7.csv", files[0][1].replace("发布日期".encode(), "旧日期字段".encode(), 1))
    errors = " ".join(parse_library_files(files).errors)
    assert "发布日期" in errors and "旧日期字段" in errors and "works_v7.csv" in errors


def test_canonical_csv_missing_table_and_manifest_are_named(populated):
    files = [(name, raw) for name, raw in unpack_csv(export_library(populated, "csv")) if name not in {"characters.csv", "manifest.json"}]
    errors = " ".join(parse_library_files(files).errors)
    assert "manifest.json" in errors and "characters.csv" in errors


def test_mixed_csv_bundle_ids_name_pair(populated):
    first = dict(unpack_csv(export_library(populated, "csv")))
    second = dict(unpack_csv(export_library(populated, "csv")))
    first["characters.csv"] = second["characters.csv"]
    errors = " ".join(parse_library_files(list(first.items())).errors)
    assert "characters.csv" in errors and "manifest.json" in errors and "导出批次不一致" in errors


def test_edited_csv_is_reviewable_and_not_silently_rejected(populated):
    files = dict(unpack_csv(export_library(populated, "csv")))
    files["works.csv"] = files["works.csv"].replace(b"baseline", b"edited")
    batch = parse_library_files(list(files.items()))
    assert not batch.errors
    assert "works.csv" in " ".join(batch.warnings)
    assert "经过编辑" in " ".join(batch.warnings)


def test_standalone_csv_duplicate_empty_table_rejected():
    files = fixtures()
    header = files[0][1].splitlines()[0] + b"\n"
    files.append(("duplicate.csv", header))
    errors = " ".join(parse_library_files(files).errors)
    assert "duplicate.csv" in errors and "works_v8.csv" in errors and "重复选择" in errors


def test_preview_does_not_write_commit_atomic_and_requires_clean_preview(tmp_path):
    store = DataStore(tmp_path / "target.db")
    batch = parse_library_files([fixtures()[0]])
    preview_library_import(store, batch)
    with pytest.raises(DataError, match="无错误"):
        commit_library_import(store, batch)
    assert store.stats()["works"] == 0
    batch = parse_library_files(fixtures())
    result = preview_library_import(store, batch)
    assert not result["errors"] and store.stats()["works"] == 0
    assert commit_library_import(store, batch)["imported"] == 8


@pytest.mark.parametrize("fmt", ["db", "csv", "xlsx"])
def test_idempotent_reimport_and_explicit_update_policy(tmp_path, populated, fmt):
    target, raw = roundtrip(tmp_path, populated, fmt)
    batch = parse_library_files([("library." + ("zip" if fmt == "csv" else fmt), raw)])
    assert not preview_library_import(target, batch)["errors"]
    revisions = {table: [(row["id"], row["revision"]) for row in target.all_records(table)] for table in TABLES}
    result = commit_library_import(target, batch, "update")
    assert result["skipped"] == 8 and result["relations_skipped"] == 1
    assert {table: [(row["id"], row["revision"]) for row in target.all_records(table)] for table in TABLES} == revisions
    preview_library_import(target, batch)
    with pytest.raises(DataError) as exc:
        commit_library_import(target, batch, "error")
    assert exc.value.status_code == 409


def test_relation_changes_since_preview_are_detected(populated):
    batch = parse_library_files([("library.db", export_library(populated, "db"))])
    assert not preview_library_import(populated, batch)["errors"]
    relation = populated.list_material_relations()[0]
    populated.update_material_relation("R001", {"description": "changed"}, relation["revision"])
    with pytest.raises(DataError, match="关联已发生变化"):
        commit_library_import(populated, batch, "update")


def test_typed_relation_failure_rolls_back_material_updates(tmp_path, populated):
    target, _ = roundtrip(tmp_path, populated, "db")
    row = populated.get_record("works", "W950")
    populated.update_record("works", row["id"], {"热度权重": "changed"}, row["revision"])
    batch = parse_library_files([("library.db", export_library(populated, "db"))])
    assert not preview_library_import(target, batch)["errors"]
    # A server-side mutation after preview still cannot leave partial writes.
    batch.relation_tables["material_relations"][0]["target_id"] = "MISSING"
    with pytest.raises(DataError):
        commit_library_import(target, batch, "update")
    assert target.get_record("works", "W950")["values"]["热度权重"] == "0"


def test_cliche_to_cliche_relation_rejected_with_id(populated):
    files = dict(unpack_csv(export_library(populated, "csv")))
    rows = list(csv.DictReader(io.StringIO(files["material_relations.csv"].decode("utf-8-sig"))))
    rows[0].update(source_table="cliches", source_id="C007", target_table="cliches", target_id="C019")
    files["material_relations.csv"] = csv_data(list(rows[0]), rows)
    errors = " ".join(parse_library_files(list(files.items())).errors)
    assert "R001" in errors and "桥段关联桥段" in errors


def test_work_cliche_bidirectional_mismatch_names_files(populated):
    files = dict(unpack_csv(export_library(populated, "csv")))
    rows = list(csv.DictReader(io.StringIO(files["works.csv"].decode("utf-8-sig"))))
    rows[0]["关联cliché"] = "[]"
    files["works.csv"] = csv_data(list(rows[0]), rows)
    errors = " ".join(parse_library_files(list(files.items())).errors)
    assert "cliches.csv" in errors and "works.csv" in errors and "W900" in errors


def test_db_read_only_import_rejects_trigger_and_unknown_tables(tmp_path, populated):
    path = tmp_path / "hostile.db"
    path.write_bytes(export_library(populated, "db"))
    with sqlite3.connect(path) as connection:
        connection.execute('CREATE TRIGGER unexpected AFTER INSERT ON works BEGIN DELETE FROM works; END')
    errors = " ".join(parse_library_files([("hostile.db", path.read_bytes())]).errors)
    assert "unexpected" in errors and "触发器" in errors


def test_db_missing_required_table_named(tmp_path, populated):
    path = tmp_path / "missing.db"
    path.write_bytes(export_library(populated, "db"))
    with sqlite3.connect(path) as connection:
        connection.execute('DROP TABLE characters')
    errors = " ".join(parse_library_files([("missing.db", path.read_bytes())]).errors)
    assert "characters.csv" in errors and "缺少必要" in errors


def test_zip_path_duplicate_and_expansion_bounded():
    for names in [["../works.csv"], ["works.csv", "works.csv"]]:
        output = io.BytesIO()
        import warnings
        with warnings.catch_warnings():
            warnings.simplefilter("ignore", UserWarning)
            with zipfile.ZipFile(output, "w") as archive:
                for name in names:
                    archive.writestr(name, "a,b\n")
        assert parse_library_files([("invalid.zip", output.getvalue())]).errors


def test_xlsx_missing_sheet_and_long_text_part_are_named(populated):
    raw = export_library(populated, "xlsx")
    book = load_workbook(io.BytesIO(raw))
    del book["characters"]
    target = io.BytesIO()
    book.save(target)
    errors = " ".join(parse_library_files([("missing.xlsx", target.getvalue())]).errors)
    assert "characters" in errors and "工作表" in errors


def test_xlsx_xml_entity_rejected(populated):
    raw = export_library(populated, "xlsx")
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(raw)) as old, zipfile.ZipFile(output, "w") as archive:
        for name in old.namelist():
            value = old.read(name)
            if name == "xl/workbook.xml":
                value = b'<!DOCTYPE workbook [<!ENTITY x "boom">]>' + value
            archive.writestr(name, value)
    errors = " ".join(parse_library_files([("entities.xlsx", output.getvalue())]).errors)
    assert "实体" in errors and "workbook.xml" in errors


def test_formats_document_all_five_tables_and_default_db():
    formats = library_formats()
    assert formats["default"] == "db"
    assert {row["table"] for row in formats["tables"]} == set(COLUMNS)
    assert {row["format"] for row in formats["formats"]} == {"db", "csv", "xlsx"}


def test_canonical_three_csv_cannot_be_downgraded_to_legacy(populated):
    files = [(name, raw) for name, raw in unpack_csv(export_library(populated, "csv")) if name in {"works.csv", "characters.csv", "cliches.csv"}]
    errors = " ".join(parse_library_files(files).errors)
    assert "manifest.json" in errors


def test_xlsx_mixed_worksheet_bundle_detected(populated):
    first = load_workbook(io.BytesIO(export_library(populated, "xlsx")))
    second = load_workbook(io.BytesIO(export_library(populated, "xlsx")))
    # Copying another export's sheet carries a different explicit bundle ID.
    bundle_column = len(COLUMNS["characters"]) + 1
    first["characters"].cell(2,bundle_column).value = second["characters"].cell(2,bundle_column).value
    output = io.BytesIO()
    first.save(output)
    errors = " ".join(parse_library_files([("mixed.xlsx", output.getvalue())]).errors)
    assert "characters" in errors and "_manifest" in errors and "不同版本" in errors


def test_xlsx_non_bmp_excel_utf16_cell_limit(tmp_path, populated):
    row = populated.get_record("works", "W950")
    text = "😀" * 20000
    populated.update_record("works", row["id"], {"剧情梗概": text}, row["revision"])
    target, raw = roundtrip(tmp_path, populated, "xlsx")
    assert target.get_record("works", "W950")["values"]["剧情梗概"] == text
    book = load_workbook(io.BytesIO(raw), read_only=True)
    assert "_long_text" in book.sheetnames
    book.close()


def test_sqlite_exchange_connections_are_explicitly_closed(populated, monkeypatch):
    import inkreference.library_exchange as exchange
    real_connect = sqlite3.connect
    tracked = []

    class TrackingConnection(sqlite3.Connection):
        explicitly_closed = False

        def close(self):
            self.explicitly_closed = True
            super().close()

    def connect(*args, **kwargs):
        if "factory" not in kwargs:
            kwargs["factory"] = TrackingConnection
            result = real_connect(*args, **kwargs)
            tracked.append(result)
            return result
        return real_connect(*args, **kwargs)

    monkeypatch.setattr(exchange.sqlite3, "connect", connect)
    raw = exchange.export_library(populated, "db")
    assert not exchange.parse_library_files([("all.db", raw)]).errors
    assert len(tracked) == 2
    assert all(connection.explicitly_closed for connection in tracked)
