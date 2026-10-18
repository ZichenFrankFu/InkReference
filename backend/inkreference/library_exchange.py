"""Portable material-library bundles, separate from full project backups.

The same five business tables and manifest travel in SQLite, CSV ZIP, or XLSX.
Parsing never executes SQL supplied by a file. All writes reuse DataStore's
transaction, revision checks, and authoritative legacy relation synchronization.
"""
from __future__ import annotations

import base64
import csv
import hashlib
import io
import json
import math
import re
import sqlite3
import tempfile
import uuid
import zipfile
from contextlib import closing
from dataclasses import dataclass, field
from datetime import date, datetime, timezone
from pathlib import Path, PurePosixPath
from typing import Any

from .csvio import ImportBatch, ImportTable, _apply_import, _csv_bytes, _decode, _snapshot
from .database import TABLES, DataError, DataStore, character_version, parse_work_relations, record_id, validate_values

FORMAT = "InkReference material library"
FORMAT_VERSION = 1
SCHEMA_VERSION = 3
MAX_BYTES = 128 * 1024 * 1024
MAX_EXPANDED_BYTES = 256 * 1024 * 1024
MAX_ROWS = 200_000
MAX_CELL_BYTES = 32 * 1024 * 1024
MAX_ZIP_FILES = 256
RELATION_COLUMNS = ["work_id", "cliche_id", "detail", "work_position", "cliche_position"]
MATERIAL_RELATION_COLUMNS = ["id", "source_table", "source_id", "target_table", "target_id", "relation_type", "description", "metadata_json"]
COLUMNS = {**{key: list(spec["columns"]) for key, spec in TABLES.items()}, "work_cliches": RELATION_COLUMNS, "material_relations": MATERIAL_RELATION_COLUMNS}
LABELS = {**{key: spec["label"] for key, spec in TABLES.items()}, "work_cliches": "作品与桥段关联", "material_relations": "素材关联说明"}
LEGACY_NAMES = {"works": "works_library_v8.csv", "characters": "character_database_v8.csv", "cliches": "cliche_database_v8.csv"}
CELL_REF = "__INKREFERENCE_TEXT__:"
BAD_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


@dataclass
class ExchangeBatch(ImportBatch):
    format: str = "csv"
    legacy: bool = False
    manifest: dict[str, Any] | None = None
    relation_tables: dict[str, list[dict[str, str]]] = field(default_factory=dict)
    filenames: dict[str, str] = field(default_factory=dict)
    relation_snapshot: str | None = None
    preview_ok: bool = False


def library_formats() -> dict[str, Any]:
    return {
        "default": "db", "schema_version": SCHEMA_VERSION,
        "tables": [{"table": key, "label": LABELS[key], "filename": f"{key}.csv", "columns": value} for key, value in COLUMNS.items()],
        "formats": [
            {"format": "db", "label": "SQLite 数据库", "extension": ".db", "description": "一个 .db 文件包含全部素材表与关联说明"},
            {"format": "csv", "label": "CSV 文件包", "extension": ".zip", "description": "ZIP 内含五张表和 manifest.json；也可同时选择这些文件"},
            {"format": "xlsx", "label": "Excel 工作簿", "extension": ".xlsx", "description": "一个工作簿包含五张表和格式信息"},
        ],
        "legacy_csv": [{"table": key, "filename": value} for key, value in LEGACY_NAMES.items()],
        "guidance": ["推荐使用 .db，整库导入导出只需一个文件。", "CSV 请上传完整 ZIP，或同时选择五张 CSV 与 manifest.json；旧 v8 格式必须同时选择作品、人物、桥段三张 CSV。", "XLSX 请保留所有工作表及 _manifest；编号、日期与长文本按文本保存。"],
    }


def _json(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, separators=(",", ":"), sort_keys=True)


def _digest(rows: list[dict[str, str]]) -> str:
    return hashlib.sha256(_json(rows).encode("utf-8")).hexdigest()


def _required(batch: ExchangeBatch) -> list[dict[str, Any]]:
    keys = TABLES if batch.legacy else COLUMNS
    return [{"table": key, "label": LABELS[key], "filename": LEGACY_NAMES[key] if batch.legacy else f"{key}.csv", "present": key in batch.filenames} for key in keys]


def _error(batch: ExchangeBatch, message: str) -> None:
    if message not in batch.errors:
        batch.errors.append(message)


def _zip_entries(data: bytes, *, xlsx: bool = False) -> dict[str, bytes]:
    if len(data) > MAX_BYTES:
        raise DataError(413, "导入文件超过 128 MB，请分库存储。")
    result: dict[str, bytes] = {}
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            infos = archive.infolist()
            if len(infos) > MAX_ZIP_FILES:
                raise DataError(422, "压缩包文件过多；请使用 InkReference 导出的完整文件包。")
            expanded = 0
            for info in infos:
                if info.is_dir():
                    continue
                name = info.filename
                path = PurePosixPath(name)
                if path.is_absolute() or ".." in path.parts or "\\" in name or info.flag_bits & 1:
                    raise DataError(422, f"压缩包成员路径或加密状态无效：{name}")
                if name in result:
                    raise DataError(422, f"压缩包包含重复文件：{name}")
                expanded += info.file_size
                if expanded > MAX_EXPANDED_BYTES or info.file_size > MAX_BYTES:
                    raise DataError(413, "压缩包解压后超过 256 MB 或单项超过 128 MB。")
                if not xlsx and len(path.parts) != 1:
                    raise DataError(422, f"CSV 文件包中 {name} 位于子目录；请将所有 CSV 与 manifest.json 放在同一层。")
                raw = archive.read(info)
                if xlsx and name.lower().endswith(".xml") and re.search(br"<!\s*(?:DOCTYPE|ENTITY)\b", raw, re.I):
                    raise DataError(422, f"Excel XML 含外部实体或文档类型定义：{name}")
                result[name] = raw
    except (zipfile.BadZipFile, RuntimeError, EOFError) as exc:
        raise DataError(422, "压缩文件不完整或格式无效，请重新导出。") from exc
    return result


def _manifest(value: Any) -> dict[str, Any]:
    if not isinstance(value, dict) or value.get("format") != FORMAT or value.get("format_version") != FORMAT_VERSION:
        raise DataError(422, "manifest.json / _manifest 不是受支持的 InkReference 素材库格式，请重新导出完整素材库。")
    if value.get("schema_version") != SCHEMA_VERSION:
        raise DataError(422, f"素材库结构版本为 {value.get('schema_version')}，当前要求 {SCHEMA_VERSION}；请使用匹配版本的应用重新导出所有表。")
    if not isinstance(value.get("bundle_id"), str) or not value["bundle_id"] or not isinstance(value.get("tables"), dict):
        raise DataError(422, "素材库格式信息缺少 bundle_id 或 tables；请重新导出完整文件包。")
    if set(value["tables"]) != set(COLUMNS):
        missing = set(COLUMNS) - set(value["tables"])
        extra = set(value["tables"]) - set(COLUMNS)
        raise DataError(422, f"格式信息中的表不完整；缺少：{'、'.join(sorted(missing)) or '无'}；未知：{'、'.join(sorted(extra)) or '无'}。")
    for table, columns in COLUMNS.items():
        info = value["tables"][table]
        if not isinstance(info, dict) or info.get("columns") != columns:
            raise DataError(422, f"{table} 的格式信息与当前字段不符，请重新导出包含全部表的同一批文件。")
    return value


def _add_table(batch: ExchangeBatch, table: str, filename: str, rows: list[dict[str, str]]) -> None:
    if table in batch.filenames:
        raise DataError(422, f"{filename} 与 {batch.filenames[table]} 都是 {LABELS[table]}表，重复选择；请保留同一版本的一份。")
    batch.filenames[table] = filename
    if table in TABLES:
        clean = [validate_values(table, row) for row in rows]
        ids: set[str] = set()
        for row in clean:
            identifier = record_id(table, row)
            if identifier in ids:
                raise DataError(422, f"{filename}：重复 ID {identifier}。")
            ids.add(identifier)
        batch.tables.append(ImportTable(table, filename, clean))
    else:
        batch.relation_tables[table] = rows


def _rows_csv(data: bytes, filename: str) -> tuple[list[str], list[dict[str, str]]]:
    text, _ = _decode(data)
    reader = csv.reader(io.StringIO(text, newline=""), strict=True)
    header = next(reader, [])
    if not header or len(header) != len(set(header)):
        raise DataError(422, f"{filename} 缺少表头或存在重复字段。")
    result = []
    total_bytes = 0
    for row in reader:
        if not row:
            continue
        if len(result) >= MAX_ROWS:
            raise DataError(413, f"{filename} 超过 {MAX_ROWS} 行。")
        if len(row) != len(header):
            raise DataError(422, f"{filename} 第 {reader.line_num} 行有 {len(row)} 项，应为 {len(header)} 项；请检查引号与换行。")
        for value in row:
            size = len(value.encode("utf-8"))
            if size > MAX_CELL_BYTES:
                raise DataError(413, f"{filename} 第 {reader.line_num} 行有字段超过 32 MB。")
            total_bytes += size
        if total_bytes > MAX_BYTES:
            raise DataError(413, f"{filename} 文本总量超过 128 MB。")
        result.append(dict(zip(header, row)))
    return header, result


def _header_error(filename: str, header: list[str]) -> str:
    plain = set(header) - {"_bundle_id", "_schema_version"}
    key = min(COLUMNS, key=lambda table: len(set(COLUMNS[table]) ^ plain))
    missing = set(COLUMNS[key]) - plain
    extra = plain - set(COLUMNS[key])
    return f"{filename} 疑似{LABELS[key]}表但字段不符；缺少：{'、'.join(sorted(missing)) or '无'}；多出：{'、'.join(sorted(extra)) or '无'}。请勿混用不同版本的表。"


def _parse_csv(batch: ExchangeBatch, files: list[tuple[str, bytes]]) -> None:
    manifests = [(name, raw) for name, raw in files if name.lower().endswith(".json")]
    if len(manifests) > 1:
        raise DataError(422, f"重复格式信息文件：{'、'.join(name for name, _ in manifests)}；请保留同一批次的 manifest.json。")
    if manifests:
        batch.manifest = _manifest(json.loads(manifests[0][1].decode("utf-8-sig")))
    batch.legacy = not bool(batch.manifest)
    versions: dict[str, str] = {}
    for filename, data in files:
        if filename.lower().endswith(".json"):
            continue
        if not filename.lower().endswith(".csv"):
            _error(batch, f"{filename} 不是 CSV；CSV 导入只应包含全部 CSV 和 manifest.json。")
            continue
        try:
            header, rows = _rows_csv(data, filename)
            if {"_bundle_id", "_schema_version"} & set(header):
                batch.legacy = False
            plain = [field for field in header if field not in {"_bundle_id", "_schema_version"}]
            matches = [table for table, columns in COLUMNS.items() if set(plain) == set(columns)]
            if len(matches) != 1:
                raise DataError(422, _header_error(filename, header))
            table = matches[0]
            if table not in TABLES:
                batch.legacy = False
            version = re.search(r"(?:^|[_\-])v(\d+)(?:[_.\-]|$)", filename, re.I)
            if version:
                versions[filename] = version.group(1)
            if batch.manifest:
                for number, row in enumerate(rows, 2):
                    if "_bundle_id" in row and row["_bundle_id"] != batch.manifest["bundle_id"]:
                        raise DataError(422, f"{filename} 第 {number} 行与 manifest.json 的导出批次不一致；{filename} 与 manifest.json 疑似来自不同版本，请重新选择同一次导出的完整文件包。")
                    if "_schema_version" in row and row["_schema_version"] != str(SCHEMA_VERSION):
                        raise DataError(422, f"{filename} 第 {number} 行结构版本与 manifest.json 不一致；请重新导出所有表。")
            _add_table(batch, table, filename, [{key: row[key] for key in COLUMNS[table]} for row in rows])
        except (DataError, csv.Error, UnicodeError) as exc:
            _error(batch, str(exc) if filename in str(exc) else f"{filename}：{exc}")
    if len(set(versions.values())) > 1:
        pairs = [f"{name}（v{version}）" for name, version in versions.items()]
        _error(batch, f"{' 与 '.join(pairs)} 的版本疑似不一致；请选用同一版本的全部文件。")
    if not batch.legacy and not batch.manifest:
        _error(batch, "缺少必要文件 manifest.json；请将 works.csv、characters.csv、cliches.csv、work_cliches.csv、material_relations.csv 与 manifest.json 一起导入，或直接选择完整 ZIP。")


def _parse_db(batch: ExchangeBatch, filename: str, data: bytes) -> None:
    if not data.startswith(b"SQLite format 3\x00"):
        raise DataError(422, f"{filename} 不是 SQLite 数据库，请使用导出的 .db 文件。")
    with tempfile.TemporaryDirectory(prefix="inkreference-import-") as directory:
        path = Path(directory) / "library.db"
        path.write_bytes(data)
        with closing(sqlite3.connect(f"{path.as_uri()}?mode=ro&immutable=1", uri=True)) as connection:
            connection.execute("PRAGMA trusted_schema = OFF")
            connection.execute("PRAGMA query_only = ON")
            connection.setlimit(sqlite3.SQLITE_LIMIT_LENGTH, MAX_CELL_BYTES)
            connection.setlimit(sqlite3.SQLITE_LIMIT_COLUMN, 128)
            remaining = [25_000]
            def limit() -> int:
                remaining[0] -= 1
                return int(remaining[0] <= 0)
            connection.set_progress_handler(limit, 1000)
            schema = connection.execute("SELECT type,name,sql FROM sqlite_schema WHERE name NOT LIKE 'sqlite_%'").fetchall()
            bad = [name for kind, name, sql in schema if kind not in {"table", "index"} or (kind == "table" and re.search(r"\bVIRTUAL\b", sql or "", re.I))]
            if bad:
                raise DataError(422, f"{filename} 包含不支持的视图、触发器或虚拟表：{'、'.join(bad)}；仅接受素材数据表。")
            names = {name for kind, name, _ in schema if kind == "table"}
            if "_manifest" not in names:
                raise DataError(422, f"{filename} 缺少 _manifest 格式表；请从 InkReference 的“导出素材库 → DB”重新导出。项目工作数据库请使用项目恢复入口。")
            if names - set(COLUMNS) - {"_manifest"}:
                raise DataError(422, f"{filename} 含未知表：{'、'.join(sorted(names - set(COLUMNS) - {'_manifest'}))}；请选择完整素材库导出文件。")
            info = connection.execute('SELECT "json" FROM "_manifest" LIMIT 2').fetchall()
            if len(info) != 1 or not isinstance(info[0][0], str):
                raise DataError(422, f"{filename} 的 _manifest 必须有一条文本格式信息。")
            batch.manifest = _manifest(json.loads(info[0][0]))
            for table, columns in COLUMNS.items():
                if table not in names:
                    continue
                found = [row[1] for row in connection.execute(f'PRAGMA table_info("{table}")')]
                if found != columns:
                    raise DataError(422, _header_error(f"{filename} / {table}", found))
                cursor = connection.execute(f'SELECT * FROM "{table}"')
                rows = []
                for row in cursor:
                    if len(rows) >= MAX_ROWS:
                        raise DataError(413, f"{filename} 的 {table} 表超过 {MAX_ROWS} 行。")
                    if any(not isinstance(value, str) for value in row):
                        raise DataError(422, f"{filename} / {table} 第 {len(rows) + 1} 行包含非文本值；请保留编号和日期的原始文本。")
                    rows.append(dict(zip(columns, row)))
                _add_table(batch, table, f"{filename} / {table}", rows)


def _excel_string(cell: Any, label: str, warnings: list[str]) -> str:
    value = cell.value
    if value is None:
        return ""
    if cell.data_type == "f":
        raise DataError(422, f"{label} 是公式；请粘贴为文本值后导入，编号和日期请设为文本格式。")
    if cell.data_type == "e":
        raise DataError(422, f"{label} 是 Excel 错误值 {value}，请修正该单元格。")
    if isinstance(value, str):
        return value
    if isinstance(value, (datetime, date)):
        raise DataError(422, f"{label} 已被 Excel 转为日期，无法还原原日期精度；请将该单元格改为文本并填写原值。")
    if isinstance(value, (int, float)) and not isinstance(value, bool) and (not isinstance(value, float) or math.isfinite(value)):
        message = f"{label} 为数字，已转为文本；如原值含前导零或长编号，请先在 Excel 中恢复原值。"
        if len(warnings) < 20:
            warnings.append(message)
        return str(value) if not isinstance(value, float) or not value.is_integer() else str(int(value))
    raise DataError(422, f"{label} 的值类型不受支持，请改为文本。")


def _parse_xlsx(batch: ExchangeBatch, filename: str, data: bytes) -> None:
    _zip_entries(data, xlsx=True)
    from openpyxl import load_workbook
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
    try:
        names = set(workbook.sheetnames)
        allowed = set(COLUMNS) | {"_manifest", "_long_text"}
        if names - allowed:
            raise DataError(422, f"{filename} 有未知工作表：{'、'.join(sorted(names - allowed))}；请保留同一素材库的全部工作表。")
        if "_manifest" not in names:
            raise DataError(422, f"{filename} 缺少必要工作表 _manifest；请使用 InkReference 导出的完整工作簿。")
        meta_rows = list(workbook["_manifest"].iter_rows(max_row=2, max_col=1))
        if not meta_rows or not isinstance(meta_rows[0][0].value, str) or (len(meta_rows) > 1 and meta_rows[1][0].value is not None):
            raise DataError(422, f"{filename} 的 _manifest 格式信息损坏。")
        batch.manifest = _manifest(json.loads(meta_rows[0][0].value))
        long_text: dict[str, list[tuple[int, str]]] = {}
        if "_long_text" in names:
            sheet = workbook["_long_text"]
            for index, row in enumerate(sheet.iter_rows(), 1):
                if index == 1:
                    if [cell.value for cell in row] != ["key", "part", "base64_utf8"]:
                        raise DataError(422, "_long_text 长文本表字段损坏；请勿修改该辅助工作表。")
                    continue
                if index > MAX_ROWS or len(row) != 3:
                    raise DataError(422, "_long_text 长文本表行数或字段数无效。")
                key, part, raw = [cell.value for cell in row]
                if not isinstance(key, str) or not isinstance(part, (str, int)) or not str(part).isdigit() or not isinstance(raw, str):
                    raise DataError(422, f"_long_text 第 {index} 行格式无效。")
                long_text.setdefault(key, []).append((int(part), raw))
        decoded: dict[str, str] = {}
        for key, parts in long_text.items():
            parts.sort()
            if [part for part, _ in parts] != list(range(len(parts))):
                raise DataError(422, f"_long_text 的 {key} 缺少或重复片段。")
            raw = "".join(value for _, value in parts)
            try:
                value = base64.b64decode(raw, validate=True).decode("utf-8", errors="strict")
            except (ValueError, UnicodeError) as exc:
                raise DataError(422, f"_long_text 的 {key} 编码损坏。") from exc
            if len(value.encode("utf-8")) > MAX_CELL_BYTES:
                raise DataError(413, f"_long_text 的 {key} 超过 32 MB。")
            decoded[key] = value
        used: set[str] = set()
        for table, columns in COLUMNS.items():
            if table not in names:
                continue
            sheet = workbook[table]
            if sheet.max_column is not None and sheet.max_column > 128:
                raise DataError(422, f"{filename} / {table} 的列数异常。")
            iterator = sheet.iter_rows()
            header_cells = next(iterator, [])
            header = [cell.value for cell in header_cells]
            expected_header = [*columns, "_bundle_id", "_schema_version"]
            if header != expected_header:
                raise DataError(422, _header_error(f"{filename} / {table}", header))
            rows = []
            for index, cells in enumerate(iterator, 2):
                if index > MAX_ROWS + 1:
                    raise DataError(413, f"{filename} / {table} 超过 {MAX_ROWS} 行。")
                if all(cell.value is None for cell in cells):
                    continue
                if len(cells) != len(expected_header):
                    raise DataError(422, f"{filename} / {table} 第 {index} 行字段数不符。")
                bundle = _excel_string(cells[-2], f"{filename} / {table}!{cells[-2].coordinate}", batch.warnings)
                schema = _excel_string(cells[-1], f"{filename} / {table}!{cells[-1].coordinate}", batch.warnings)
                if bundle != batch.manifest["bundle_id"] or schema != str(SCHEMA_VERSION):
                    raise DataError(422, f"{filename} / {table} 第 {index} 行与 _manifest 的批次或结构版本不符；{table} 与 _manifest 工作表疑似来自不同版本，请重新选择同次导出的完整工作簿。")
                values = {}
                for column, cell in zip(columns, cells):
                    value = _excel_string(cell, f"{filename} / {table}!{cell.coordinate}", batch.warnings)
                    if value.startswith(CELL_REF):
                        key = value[len(CELL_REF):]
                        if key not in decoded or key != f"{table}/{index - 2}/{columns.index(column)}":
                            raise DataError(422, f"{filename} / {table}!{cell.coordinate} 的长文本片段缺失或位置不符，请恢复 _long_text 工作表。")
                        used.add(key)
                        value = decoded[key]
                    values[column] = value
                rows.append(values)
            _add_table(batch, table, f"{filename} / {table}", rows)
        if set(decoded) - used:
            batch.warnings.append("部分长文本辅助片段已不再被使用；对应正文单元格可能经过手动编辑，请在预览中核对。")
    finally:
        workbook.close()


def _validate_complete(batch: ExchangeBatch) -> None:
    for item in _required(batch):
        if not item["present"]:
            _error(batch, f"缺少必要{'工作表' if batch.format == 'xlsx' else '表/文件'}：{item['label']}（{item['filename']}）。请选用同一批次的全部表，空表也应保留表头。")
    if not batch.manifest and not batch.legacy:
        return
    rows_by_table = {item.table: item.rows for item in batch.tables}
    changed_tables: set[str] = set()
    if batch.manifest:
        for table, rows in {**rows_by_table, **batch.relation_tables}.items():
            info = batch.manifest["tables"][table]
            if info.get("sha256") != _digest(rows) or info.get("rows") != len(rows):
                changed_tables.add(table)
    indexes = {table: {record_id(table, row): row for row in rows_by_table.get(table, [])} for table in TABLES}
    # Validate all incoming references independently of local records: a complete
    # bundle must not accidentally borrow a missing parent from the current DB.
    if set(TABLES) <= set(batch.filenames):
        for child in rows_by_table.get("characters", []):
            parent = indexes["works"].get(child["作品ID"])
            pair = f"{batch.filenames['characters']} 与 {batch.filenames['works']}"
            if not parent:
                _error(batch, f"{pair} 疑似版本不一致：人物 {child['人物ID']} 引用作品 {child['作品ID']}，作品表缺少该 ID。")
            elif child["作品"] and child["作品"] != parent["作品"]:
                _error(batch, f"{pair} 疑似版本不一致：人物 {child['人物ID']} 的作品名称“{child['作品']}”与作品 {child['作品ID']} 的“{parent['作品']}”不符。")
            elif (batch.legacy or changed_tables & {"works", "characters"}) and child["版本"] and parent["版本与原作关系"] and character_version(parent["版本与原作关系"]) != child["版本"]:
                # Edition strings are intentionally lossless and can differ in
                # provenance suffix. Only flag for review, do not rewrite them.
                batch.warnings.append(f"{pair}：人物 {child['人物ID']} 的版本“{child['版本']}”与作品版本“{parent['版本与原作关系']}”不同，请核对是否为同一版。")
        for wid, work in indexes["works"].items():
            for cid in parse_work_relations(work["典型桥段（剧透）"])[1]:
                if cid not in indexes["cliches"]:
                    _error(batch, f"{batch.filenames['works']} 与 {batch.filenames['cliches']} 疑似版本不一致：作品 {wid} 引用桥段 {cid}，桥段表缺少该 ID。")
    if batch.manifest:
        all_rows = {**rows_by_table, **batch.relation_tables}
        for table, rows in all_rows.items():
            info = batch.manifest["tables"][table]
            if info.get("sha256") != _digest(rows) or info.get("rows") != len(rows):
                batch.warnings.append(f"{batch.filenames[table]} 的内容或行数与导出时不同，可能经过编辑；请核对预览中的关联与版本提示。")
        _validate_relation_rows(batch, indexes)


def _validate_relation_rows(batch: ExchangeBatch, indexes: dict[str, dict[str, dict[str, str]]]) -> None:
    seen: set[tuple[str, str]] = set()
    detail_map = {(wid, cid): detail for wid, work in indexes["works"].items() for cid, detail in parse_work_relations(work["典型桥段（剧透）"])[1].items()}
    for row in batch.relation_tables.get("work_cliches", []):
        key = row["work_id"], row["cliche_id"]
        label = batch.filenames.get("work_cliches", "work_cliches.csv")
        if key in seen:
            _error(batch, f"{label} 重复关联：作品 {key[0]} / 桥段 {key[1]}。")
        seen.add(key)
        for table, identifier in zip(("works", "cliches"), key):
            if identifier not in indexes[table]:
                _error(batch, f"{label} 与 {batch.filenames.get(table, table)} 疑似版本不一致：关联引用的 {LABELS[table]} ID {identifier} 不存在。")
        for field_name in ("work_position", "cliche_position"):
            if not re.fullmatch(r"\d{1,10}", row[field_name]):
                _error(batch, f"{label} 的 {key[0]} / {key[1]} 字段 {field_name} 应为非负整数。")
        if key not in detail_map or detail_map[key] != row["detail"]:
            _error(batch, f"{label} 与 {batch.filenames.get('works', 'works.csv')} 对作品 {key[0]} / 桥段 {key[1]} 的关联说明不一致；请检查是否混用版本。")
    if seen != set(detail_map):
        for wid, cid in sorted(set(detail_map) - seen):
            _error(batch, f"work_cliches.csv 缺少 works.csv 中的关联：作品 {wid} / 桥段 {cid}；这两个文件疑似版本不一致。")
    relation_ids: set[str] = set()
    for row in batch.relation_tables.get("material_relations", []):
        rid = row["id"]
        if not rid or rid in relation_ids:
            _error(batch, f"material_relations.csv 存在空或重复关联 ID：{rid}。")
        relation_ids.add(rid)
        for prefix in ("source", "target"):
            table, identifier = row[f"{prefix}_table"], row[f"{prefix}_id"]
            if table not in TABLES or identifier not in indexes.get(table, {}):
                _error(batch, f"material_relations.csv 与 {batch.filenames.get(table, table + '.csv')} 疑似版本不一致：关联 {rid} 引用 {table}/{identifier}，对应记录不存在。")
        if row["source_table"] == row["target_table"] == "cliches":
            _error(batch, f"material_relations.csv 关联 {rid} 为桥段关联桥段；请移除该关联。")
        if not row["relation_type"].strip():
            _error(batch, f"material_relations.csv 关联 {rid} 缺少 relation_type 关联含义。")
        try:
            if not isinstance(json.loads(row["metadata_json"]), dict):
                raise ValueError()
        except (ValueError, TypeError):
            _error(batch, f"material_relations.csv 关联 {rid} 的 metadata_json 应为 JSON 对象。")


def parse_library_files(files: list[tuple[str, bytes]]) -> ExchangeBatch:
    batch = ExchangeBatch()
    if not files:
        batch.errors.append("请选择一个 .db、完整 CSV 文件包或 .xlsx 工作簿。")
        return batch
    try:
        if len(files) > 32 or sum(len(data) for _, data in files) > MAX_BYTES:
            raise DataError(413, "一次最多选择 32 个文件，总大小不能超过 128 MB。")
        suffixes = {Path(name).suffix.lower() for name, _ in files}
        if suffixes <= {".csv", ".json"}:
            _parse_csv(batch, files)
        elif len(files) != 1:
            raise DataError(422, "请导入一个 .db、一个 .xlsx、一个 CSV ZIP；仅 CSV 模式支持同时选择多张表，不能混用文件格式。")
        else:
            filename, data = files[0]
            suffix = Path(filename).suffix.lower()
            if suffix in {".db", ".sqlite", ".sqlite3"}:
                batch.format = "db"
                _parse_db(batch, filename, data)
            elif suffix == ".xlsx":
                batch.format = "xlsx"
                _parse_xlsx(batch, filename, data)
            elif suffix == ".zip":
                _parse_csv(batch, list(_zip_entries(data).items()))
            else:
                raise DataError(422, f"{filename} 格式不支持；请选择 .db、完整 CSV 文件包或 .xlsx。")
        _validate_complete(batch)
    except (DataError, csv.Error, UnicodeError, ValueError, KeyError, sqlite3.Error) as exc:
        _error(batch, str(exc))
    except Exception as exc:
        # OOXML parsers raise several package-specific errors. Do not expose
        # local paths or traceback contents originating in an untrusted file.
        _error(batch, f"文件格式损坏或不受支持（{type(exc).__name__}），请重新导出完整素材库。")
    return batch


def _relation_snapshot(connection: sqlite3.Connection) -> str:
    relations = {}
    names = {row[0] for row in connection.execute("SELECT name FROM sqlite_schema WHERE type='table'")}
    for table in ("work_cliches", "material_relations"):
        if table in names:
            relations[table] = [dict(row) for row in connection.execute(f'SELECT * FROM "{table}" ORDER BY 1,2')]
    return hashlib.sha256(_json(relations).encode("utf-8")).hexdigest()


def _apply(store: DataStore, connection: sqlite3.Connection, batch: ExchangeBatch, policy: str) -> dict[str, int]:
    previous_edges = {(row[0], row[1]) for row in connection.execute("SELECT work_id,cliche_id FROM work_cliches")}
    result = _apply_import(store, connection, batch, policy)
    relation_counts = {"relations_imported": 0, "relations_updated": 0, "relations_skipped": 0}
    for row in batch.relation_tables.get("work_cliches", []):
        key = row["work_id"], row["cliche_id"]
        current = connection.execute("SELECT * FROM work_cliches WHERE work_id=? AND cliche_id=?", key).fetchone()
        # The legacy synchronization has already validated and created edges;
        # restore both independently ordered positions from the portable bundle.
        expected = (row["detail"], int(row["work_position"]), int(row["cliche_position"]))
        if current is None:
            raise DataError(422, f"work_cliches.csv 的作品 {key[0]} / 桥段 {key[1]} 与作品表不一致；未导入。")
        old = (current["detail"], current["work_position"], current["cliche_position"])
        if old != expected and (policy == "update" or key not in previous_edges):
            connection.execute("UPDATE work_cliches SET detail=?,work_position=?,cliche_position=? WHERE work_id=? AND cliche_id=?", (*expected, *key))
            store._bump(connection, "works", {key[0]})
            store._bump(connection, "cliches", {key[1]})
    for row in batch.relation_tables.get("material_relations", []):
        existing = connection.execute("SELECT * FROM material_relations WHERE id=?", (row["id"],)).fetchone()
        same = existing is not None and all(str(existing[key]) == row[key] for key in MATERIAL_RELATION_COLUMNS)
        if same or (existing is not None and policy == "skip"):
            relation_counts["relations_skipped"] += 1
            continue
        if existing is not None and policy == "error":
            raise DataError(409, f"素材关联 ID 已存在：{row['id']}；请选择跳过或显式更新。")
        store._upsert_material_relation(connection, row, conflict_policy=policy)
        relation_counts["relations_updated" if existing is not None else "relations_imported"] += 1
    return {**result, **relation_counts}


def preview_library_import(store: DataStore, batch: ExchangeBatch) -> dict[str, Any]:
    errors = list(batch.errors)
    tables = []
    with closing(sqlite3.connect(":memory:")) as simulation:
        simulation.row_factory = sqlite3.Row
        simulation.execute("PRAGMA foreign_keys=ON")
        with store.connect() as source:
            source.execute("BEGIN")
            batch.snapshot = _snapshot(source)
            batch.relation_snapshot = _relation_snapshot(source)
            stats = store._stats(source)
            for item in batch.tables:
                existing = sum(source.execute(f"SELECT 1 FROM {item.table} WHERE id=?", (record_id(item.table, row),)).fetchone() is not None for row in item.rows)
                tables.append({"table": item.table, "filename": item.filename, "rows": len(item.rows), "new": len(item.rows) - existing, "existing": existing})
            for table, rows in batch.relation_tables.items():
                if table == "material_relations":
                    existing = sum(source.execute("SELECT 1 FROM material_relations WHERE id=?", (row["id"],)).fetchone() is not None for row in rows)
                else:
                    existing = sum(source.execute("SELECT 1 FROM work_cliches WHERE work_id=? AND cliche_id=?", (row["work_id"], row["cliche_id"])).fetchone() is not None for row in rows)
                tables.append({"table": table, "filename": batch.filenames[table], "rows": len(rows), "new": len(rows) - existing, "existing": existing})
            source.backup(simulation)
        try:
            if not errors:
                simulation.execute("BEGIN IMMEDIATE")
                _apply(store, simulation, batch, "update")
                stats = store._stats(simulation)
        except (DataError, sqlite3.Error) as exc:
            message = str(exc)
            # Name the participating files when the lower-level v8 validator found
            # edition labels or inconsistent reverse associations.
            if "出现作品" in message or "双向桥段" in message:
                message = f"{batch.filenames.get('cliches', 'cliches.csv')} 与 {batch.filenames.get('works', 'works.csv')} 疑似版本不一致：{message}。请核对作品 ID、名称、版本与桥段关联。"
            errors.append(message)
        batch.preview_ok = not errors
        return {"tables": tables, "errors": errors, "warnings": list(dict.fromkeys(batch.warnings)), "stats": stats,
                "format": batch.format, "schema_version": SCHEMA_VERSION, "bundle_id": batch.manifest["bundle_id"] if batch.manifest else None,
                "legacy": batch.legacy, "required_tables": _required(batch), "guidance": library_formats()["guidance"]}


def commit_library_import(store: DataStore, batch: ExchangeBatch, policy: str = "skip") -> dict[str, Any]:
    if not batch.preview_ok:
        raise DataError(409, "请先完成无错误的素材库导入预览，再确认导入。")
    with store.transaction() as connection:
        if batch.snapshot != _snapshot(connection) or batch.relation_snapshot != _relation_snapshot(connection):
            raise DataError(409, "预览后素材或关联已发生变化，请重新预览完整素材库后导入。")
        try:
            counts = _apply(store, connection, batch, policy)
        except sqlite3.IntegrityError as exc:
            raise DataError(422, "导入违反素材或关联约束，未写入任何记录；请核对完整文件包中的引用 ID。") from exc
        return {**counts, "stats": store._stats(connection)}


def _collect(store: DataStore) -> tuple[dict[str, list[dict[str, str]]], dict[str, Any]]:
    with store.connect() as connection:
        connection.execute("BEGIN")
        rows = {table: [item["values"] for item in store._all(connection, table)] for table in TABLES}
        rows["work_cliches"] = [{column: str(row[column]) for column in RELATION_COLUMNS} for row in connection.execute("SELECT * FROM work_cliches ORDER BY work_id,work_position,cliche_id")]
        rows["material_relations"] = [{column: str(row[column]) for column in MATERIAL_RELATION_COLUMNS} for row in connection.execute("SELECT * FROM material_relations ORDER BY id")]
    total_text_bytes = 0
    for table, values in rows.items():
        if len(values) > MAX_ROWS:
            raise DataError(413, f"{table} 超过 {MAX_ROWS} 行，请拆分素材库后导出。")
        for row in values:
            for value in row.values():
                size = len(value.encode("utf-8"))
                if size > MAX_CELL_BYTES:
                    raise DataError(413, f"{table} 有字段超过 32 MB，请拆分长文本后导出。")
                total_text_bytes += size
        if total_text_bytes > MAX_BYTES:
            raise DataError(413, "素材文本总量超过 128 MB，请拆分素材库后导出。")
    manifest = {"format": FORMAT, "format_version": FORMAT_VERSION, "schema_version": SCHEMA_VERSION, "bundle_id": uuid.uuid4().hex,
                "created_at": datetime.now(timezone.utc).isoformat(), "encoding": "UTF-8", "tables": {table: {"columns": columns, "rows": len(rows[table]), "sha256": _digest(rows[table])} for table, columns in COLUMNS.items()}}
    return rows, manifest


def _export_db(rows: dict[str, list[dict[str, str]]], manifest: dict[str, Any]) -> bytes:
    with closing(sqlite3.connect(":memory:")) as connection:
        connection.execute('CREATE TABLE "_manifest" ("json" TEXT NOT NULL)')
        connection.execute('INSERT INTO "_manifest" VALUES (?)', (_json(manifest),))
        for table, columns in COLUMNS.items():
            declaration = ",".join('"' + column.replace('"', '""') + '" TEXT NOT NULL' for column in columns)
            connection.execute(f'CREATE TABLE "{table}" ({declaration})')
            placeholders = ",".join("?" for _ in columns)
            connection.executemany(f'INSERT INTO "{table}" VALUES ({placeholders})', [[row[column] for column in columns] for row in rows[table]])
        connection.commit()
        return connection.serialize()


def _export_csv(rows: dict[str, list[dict[str, str]]], manifest: dict[str, Any]) -> bytes:
    target = io.BytesIO()
    with zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as archive:
        for table, columns in COLUMNS.items():
            extended = [{**row, "_bundle_id": manifest["bundle_id"], "_schema_version": str(SCHEMA_VERSION)} for row in rows[table]]
            archive.writestr(f"{table}.csv", _csv_bytes([*columns, "_bundle_id", "_schema_version"], extended))
        archive.writestr("manifest.json", _json(manifest).encode("utf-8"))
    return target.getvalue()


def _export_xlsx(rows: dict[str, list[dict[str, str]]], manifest: dict[str, Any]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font, PatternFill
    workbook = Workbook(write_only=True)
    long_rows: list[list[str]] = []
    for table, columns in COLUMNS.items():
        sheet = workbook.create_sheet(table)
        sheet.freeze_panes = "A2"
        header = []
        for column in [*columns, "_bundle_id", "_schema_version"]:
            cell = WriteOnlyCell(sheet, value=column)
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = PatternFill("solid", fgColor="2563EB")
            header.append(cell)
        sheet.append(header)
        for index, row in enumerate(rows[table]):
            cells = []
            for column_index, column in enumerate(columns):
                value = row[column]
                if len(value.encode("utf-16-le")) // 2 > 32767 or BAD_XML.search(value) or "\r" in value or re.search(r"_x[0-9a-fA-F]{4}_", value) or value.startswith(CELL_REF):
                    key = f"{table}/{index}/{column_index}"
                    encoded = base64.b64encode(value.encode("utf-8")).decode("ascii")
                    for part, start in enumerate(range(0, len(encoded), 30000)):
                        long_rows.append([key, str(part), encoded[start:start + 30000]])
                    # Empty literal beginning in marker is never ambiguous.
                    if not encoded:
                        long_rows.append([key, "0", ""])
                    value = CELL_REF + key
                cell = WriteOnlyCell(sheet, value=value)
                cell.data_type = "s"  # An '=' prefix remains text, never a formula.
                cell.number_format = "@"
                cells.append(cell)
            for value in (manifest["bundle_id"], str(SCHEMA_VERSION)):
                cell = WriteOnlyCell(sheet, value=value)
                cell.data_type = "s"
                cells.append(cell)
            sheet.append(cells)
    meta = workbook.create_sheet("_manifest")
    meta.append([_json(manifest)])
    meta.sheet_state = "hidden"
    if long_rows:
        long_sheet = workbook.create_sheet("_long_text")
        long_sheet.append(["key", "part", "base64_utf8"])
        for row in long_rows:
            long_sheet.append(row)
        long_sheet.sheet_state = "hidden"
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def export_library(store: DataStore, format: str = "db") -> bytes:
    if format not in {"db", "csv", "xlsx"}:
        raise DataError(422, "导出格式必须为 db、csv 或 xlsx。")
    rows, manifest = _collect(store)
    if format == "csv" and any("\x00" in value for values in rows.values() for row in values for value in row.values()):
        raise DataError(422, "素材包含 NUL 控制字符，CSV 无法可靠表示；请使用 DB 或 XLSX 导出以保留原文。")
    data = {"db": _export_db, "csv": _export_csv, "xlsx": _export_xlsx}[format](rows, manifest)
    if len(data) > MAX_BYTES:
        raise DataError(413, "导出文件超过 128 MB，请拆分素材库后导出。")
    return data
